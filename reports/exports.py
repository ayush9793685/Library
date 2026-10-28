"""Shared Excel / PDF export used by every report (fees and attendance).

All headings are English. A report passes headers + rows (+ optional totals row) and
gets back the right HttpResponse for ?export=xlsx or ?export=pdf.
"""
import io
from datetime import date, datetime

from django.http import HttpResponse
from django.utils import timezone

_FORMULA_CHARS = ("=", "+", "-", "@", "\t", "\r")


def _cell(value):
    """Plain value for spreadsheets/PDF; guards against formulas typed into names or references."""
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value
    if isinstance(value, str) and value.startswith(_FORMULA_CHARS):
        return "'" + value
    return value


def _text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return timezone.localtime(value).strftime("%d %b %Y %H:%M") if timezone.is_aware(value) else value.strftime("%d %b %Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d %b %Y")
    return str(value)


def xlsx_response(filename, title, subtitle, headers, rows, totals=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Report")[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle or ""])
    ws.append([f"Generated {timezone.localtime().strftime('%d %b %Y %H:%M')}"])
    ws.append([])
    ws.append(list(headers))
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append([_cell(v) for v in row])
    if totals:
        ws.append([_cell(v) for v in totals])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for idx, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(_text(r[idx - 1])) for r in rows if idx - 1 < len(r)] + [8])
        ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 40)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response


def pdf_response(filename, title, subtitle, headers, rows, totals=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
        topMargin=10 * mm, bottomMargin=12 * mm, title=title,
    )
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=7.5, leading=9)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    bold = ParagraphStyle("bold", parent=cell, fontName="Helvetica-Bold")

    def para(value, style):
        return Paragraph(escape(_text(value)), style)

    data = [[para(h, head) for h in headers]]
    data += [[para(v, cell) for v in row] for row in rows]
    if totals:
        data.append([para(v, bold) for v in totals])
    table = Table(data, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1d4ed8")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
    ]
    if totals:
        style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e2e8f0")))
    table.setStyle(TableStyle(style))
    story = [
        Paragraph(escape(title), styles["Title"]),
        Paragraph(escape(subtitle or ""), styles["Normal"]),
        Paragraph(f"Generated {timezone.localtime().strftime('%d %b %Y %H:%M')}", styles["Italic"]),
        Spacer(1, 4 * mm),
        table,
    ]
    doc.build(story)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
    return response


def maybe_export(request, *, filename, title, subtitle, headers, rows, totals=None):
    """Return an Excel/PDF response when ?export= asks for one, else None."""
    kind = request.GET.get("export", "")
    stamp = timezone.localdate().strftime("%Y%m%d")
    if kind == "xlsx":
        return xlsx_response(f"{filename}_{stamp}", title, subtitle, headers, rows, totals)
    if kind == "pdf":
        return pdf_response(f"{filename}_{stamp}", title, subtitle, headers, rows, totals)
    return None


def export_querystring(request):
    """The page's current filters as a query string without the export flag."""
    params = request.GET.copy()
    params.pop("export", None)
    params.pop("page", None)
    return params.urlencode()
