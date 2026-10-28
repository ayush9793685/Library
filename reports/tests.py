import io
import tempfile
from datetime import timedelta

from django.test import TestCase, override_settings
from django.utils import timezone

from accounts.models import User
from attendance.models import StaffAttendance, StudentAttendance
from students.testutils import batch, create_student, make_admin, unique_aadhaar

PRIVATE = tempfile.mkdtemp()


@override_settings(PRIVATE_MEDIA_ROOT=PRIVATE)
class ReportTests(TestCase):
    def setUp(self):
        self.admin = make_admin()
        self.client.force_login(self.admin)
        self.number = unique_aadhaar()
        self.s1 = create_student(self.client, batch("Pratibha"), aadhaar_number=self.number)
        self.s2 = create_student(self.client, batch("Sankalp"), payment_mode="upi", payment_reference="UTR1")
        self.s3 = create_student(self.client, batch("Pragati"), collect_fee="")  # no payment yet
        self.today = timezone.localdate()

    def _xlsx(self, response):
        from openpyxl import load_workbook
        self.assertIn("spreadsheetml", response["Content-Type"])
        return load_workbook(io.BytesIO(response.content)).active

    def test_date_wise_default_is_current_month_with_cash_upi_totals(self):
        r = self.client.get("/reports/fees/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context["totals"][1], 2)           # receipts
        self.assertEqual(r.context["totals"][2], "500.00")    # cash (Pratibha)
        self.assertEqual(r.context["totals"][3], "1,000.00")  # UPI (Sankalp)
        self.assertEqual(r.context["totals"][4], "1,500.00")

    def test_range_overrides_month(self):
        last_year = self.today - timedelta(days=400)
        r = self.client.get("/reports/fees/", {"date_from": last_year.isoformat(), "date_to": (last_year + timedelta(days=5)).isoformat()})
        self.assertEqual(r.context["totals"][4], "0.00")
        r = self.client.get("/reports/fees/", {"date_from": self.today.isoformat(), "date_to": self.today.isoformat()})
        self.assertEqual(r.context["totals"][4], "1,500.00")

    def test_batch_wise_total_paid_balance(self):
        r = self.client.get("/reports/fees/", {"view": "batch"})
        rows = {row[0]: row for row in r.context["rows"]}
        self.assertEqual(rows["Pratibha"][3], 1)
        self.assertEqual(rows["Pratibha"][4], "500.00")   # total fees
        self.assertEqual(rows["Pratibha"][5], "500.00")   # paid
        self.assertEqual(rows["Pratibha"][6], "0.00")     # balance
        self.assertEqual(rows["Pragati"][5], "0.00")
        self.assertEqual(rows["Pragati"][6], "600.00")    # unpaid student owes the batch fee
        self.assertEqual(r.context["totals"][4], "2,100.00")
        self.assertEqual(r.context["totals"][5], "1,500.00")
        self.assertEqual(r.context["totals"][6], "600.00")

    def test_student_wise_and_filters(self):
        r = self.client.get("/reports/fees/", {"view": "student"})
        self.assertEqual(len(r.context["rows"]), 3)
        r = self.client.get("/reports/fees/", {"view": "student", "student": self.s2.student_id})
        self.assertEqual(len(r.context["rows"]), 1)
        r = self.client.get("/reports/fees/", {"view": "receipts", "mode": "upi"})
        self.assertEqual(len(r.context["rows"]), 1)
        r = self.client.get("/reports/fees/", {"view": "receipts", "batch": batch("Pratibha").pk})
        self.assertEqual(len(r.context["rows"]), 1)

    def test_cancelled_receipts_are_excluded(self):
        from fees.services import cancel_payment
        cancel_payment(self.s1.payments.get(), reason="test", user=self.admin)
        r = self.client.get("/reports/fees/")
        self.assertEqual(r.context["totals"][4], "1,000.00")

    def test_every_fees_view_downloads_excel_and_pdf(self):
        for view in ["date", "batch", "student", "receipts"]:
            x = self.client.get("/reports/fees/", {"view": view, "export": "xlsx"})
            self.assertEqual(x.status_code, 200, view)
            ws = self._xlsx(x)
            self.assertGreaterEqual(ws.max_row, 5)
            p = self.client.get("/reports/fees/", {"view": view, "export": "pdf"})
            self.assertEqual(p.status_code, 200, view)
            self.assertEqual(p["Content-Type"], "application/pdf")
            self.assertEqual(p.content[:4], b"%PDF")

    def test_excel_content_matches_screen(self):
        ws = self._xlsx(self.client.get("/reports/fees/", {"view": "receipts", "export": "xlsx"}))
        values = [[c.value for c in row] for row in ws.iter_rows(min_row=5)]
        self.assertEqual(values[0][0], "Receipt")
        self.assertEqual(values[-1][-1], "1,500.00")

    def test_report_pages_have_all_three_buttons(self):
        for url in ["/reports/fees/", "/attendance/reports/students/", "/attendance/reports/staff/"]:
            html = self.client.get(url).content.decode()
            self.assertIn("Download Excel", html, url)
            self.assertIn("Download PDF", html, url)
            self.assertIn("Print Report", html, url)

    def test_formula_injection_is_neutralised_in_excel(self):
        self.s1.first_name = "=HYPERLINK(\"http://evil\")"
        self.s1.save()
        ws = self._xlsx(self.client.get("/reports/fees/", {"view": "receipts", "export": "xlsx"}))
        names = [row[3].value for row in ws.iter_rows(min_row=6)]
        self.assertTrue(any(n and n.startswith("'=") for n in names))

    def test_student_attendance_reports(self):
        StudentAttendance.objects.create(student=self.s1, date=self.today, status="present")
        StudentAttendance.objects.create(student=self.s2, date=self.today, status="absent")
        StudentAttendance.objects.create(student=self.s1, date=self.today - timedelta(days=1), status="leave")
        r = self.client.get("/attendance/reports/students/", {"date_from": (self.today - timedelta(days=1)).isoformat(), "date_to": self.today.isoformat()})
        self.assertEqual(len(r.context["rows"]), 3)
        r = self.client.get("/attendance/reports/students/", {"view": "summary", "month": f"{self.today:%Y-%m}"})
        totals = r.context["totals"]
        self.assertEqual(totals[4:7], [1, 1, 0] if (self.today - timedelta(days=1)).month != self.today.month else [1, 1, 1])
        r = self.client.get("/attendance/reports/students/", {"batch": batch("Pratibha").pk, "status": "present"})
        self.assertEqual(len(r.context["rows"]), 1)
        for view in ["records", "summary"]:
            for kind, magic in [("xlsx", b"PK"), ("pdf", b"%PDF")]:
                resp = self.client.get("/attendance/reports/students/", {"view": view, "export": kind})
                self.assertEqual(resp.status_code, 200)
                self.assertEqual(resp.content[:len(magic)], magic)

    def test_staff_attendance_reports_export(self):
        StaffAttendance.objects.create(user=self.admin, date=self.today, status="present")
        for view in ["records", "summary"]:
            for kind, magic in [("xlsx", b"PK"), ("pdf", b"%PDF")]:
                resp = self.client.get("/attendance/reports/staff/", {"view": view, "export": kind})
                self.assertEqual(resp.status_code, 200)
                self.assertEqual(resp.content[:len(magic)], magic)

    def test_students_csv_masks_aadhaar_and_has_seat(self):
        body = self.client.get("/reports/students/csv/").content.decode("utf-8-sig")
        self.assertNotIn(self.number, body)
        self.assertIn(f"XXXX XXXX {self.number[-4:]}", body)
        self.assertIn("Seat No", body)

    def test_reports_are_admin_only(self):
        self.client.logout()
        stu = User.objects.create_user(username="plain", password="StrongPass@123", role="student")
        self.client.force_login(stu)
        for url in ["/reports/", "/reports/fees/", "/attendance/reports/students/", "/attendance/reports/staff/",
                    "/reports/students/csv/", "/reports/fees/?export=xlsx"]:
            self.assertEqual(self.client.get(url).status_code, 403, url)
